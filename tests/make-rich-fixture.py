# Rich synthetic workbook for the format-safe save tests (no real data).
# Starts from SAMPLE_Lenovo_Inventory_SYNTH.xlsx and adds the Excel features
# SheetJS CE drops on a full rewrite: cell styles, number formats, column widths,
# freeze panes, conditional formatting, data validation, an Excel table,
# a chart, a formula sheet, a defined name, merged cells, a comment, and a
# hyperlink containing "&". Also writes an .xlsm twin with a vbaProject.bin.
# Usage: python3 tests/make-rich-fixture.py
import os, zipfile, shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "SAMPLE_Lenovo_Inventory_SYNTH.xlsx")
OUT = os.path.join(HERE, "SAMPLE_Lenovo_Inventory_RICH.xlsx")
OUTM = os.path.join(HERE, "SAMPLE_Lenovo_Inventory_RICH.xlsm")
VBA_SRC = os.path.join(HERE, "SAMPLE_Lenovo_Inventory_SYNTH.xlsm")

wb = openpyxl.load_workbook(SRC)
HDR = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="0F6CBD"),
           alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
thin = Side(style="thin", color="BDBDBD")

for ws in wb:
    for c in ws[1]:
        if c.value is not None:
            c.font, c.fill, c.alignment = HDR["font"], HDR["fill"], HDR["alignment"]
            c.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18

dev = wb["02 DEVICES"]
hdr = [c.value for c in dev[1]]
for name in ("Date Delivered", "Warranty Start", "Warranty End"):
    if name in hdr:
        col = hdr.index(name) + 1
        for r in range(2, dev.max_row + 1):
            dev.cell(r, col).number_format = "dd-mmm-yyyy"
cond = openpyxl.utils.get_column_letter(hdr.index("Condition") + 1)
last = dev.max_row
dev.conditional_formatting.add(f"{cond}2:{cond}{last}",
    FormulaRule(formula=[f'ISNUMBER(SEARCH("repair",{cond}2))'], fill=PatternFill("solid", fgColor="FDE7E9"), font=Font(color="B10E1C")))
dv = DataValidation(type="list", formula1='"Good,For Repair,Retired,Disposed"', allow_blank=True)
dev.add_data_validation(dv); dv.add(f"{cond}2:{cond}{last}")
lastcol = openpyxl.utils.get_column_letter(dev.max_column)
t = Table(displayName="tblDevices", ref=f"A1:{lastcol}{last}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
dev.add_table(t)
dev["A2"].comment = Comment("Synthetic sample row", "fixture")

tix = wb["03 TECH SUPPORT LOGS"]
th = [c.value for c in tix[1]]
st = openpyxl.utils.get_column_letter(th.index("Status") + 1)
tix.conditional_formatting.add(f"{st}2:{st}{tix.max_row}",
    CellIsRule(operator="equal", formula=['"Open"'], fill=PatternFill("solid", fgColor="FFF4CE")))

po = wb["07 PURCHASE ORDER"]
ph = [c.value for c in po[1]]
lc = ph.index("LINK") + 1
po.cell(2, lc).value = "PO link"
po.cell(2, lc).hyperlink = "https://example.sharepoint.com/sites/x/Doc.aspx?id=1&web=1&e=abc"

s = wb.create_sheet("SUMMARY", 0)
s.merge_cells("A1:D1"); s["A1"] = "Inventory summary (synthetic)"; s["A1"].font = Font(bold=True, size=14)
s["A3"], s["B3"] = "Devices", f"=COUNTA('02 DEVICES'!A2:A{last})"
s["A4"], s["B4"] = "Tickets", f"=COUNTA('03 TECH SUPPORT LOGS'!A2:A{tix.max_row})"
s["A5"], s["B5"] = "Deployments", f"=COUNTA('04 RAKSO INV.'!A2:A{wb['04 RAKSO INV.'].max_row})"
for r in (3, 4, 5): s.cell(r, 2).number_format = "#,##0"
ch = BarChart(); ch.title = "Inventory counts"
ch.add_data(Reference(s, min_col=2, min_row=3, max_row=5)); ch.set_categories(Reference(s, min_col=1, min_row=3, max_row=5))
s.add_chart(ch, "D3")
wb.defined_names["DeviceSerials"] = DefinedName("DeviceSerials", attr_text=f"'02 DEVICES'!$A$2:$A${last}")
wb.save(OUT)

# .xlsm twin: same parts + vbaProject.bin, macro-enabled main content type
with zipfile.ZipFile(VBA_SRC) as z: vba = z.read("xl/vbaProject.bin")
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(OUTM, "w", zipfile.ZIP_DEFLATED) as zout:
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == "[Content_Types].xml":
            data = data.replace(b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                                b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            data = data.replace(b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
        if it.filename == "xl/_rels/workbook.xml.rels":
            data = data.replace(b"</Relationships>", b'<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
        zout.writestr(it, data)
    zout.writestr("xl/vbaProject.bin", vba)
print("wrote", OUT, "and", OUTM)

# --- Audit fixture (v1.28 Audit editor): SAMPLE_AUDIT_RICH.xlsx -------------
# Step 1 (node): tests/make-audit-fixture.cjs writes 12 month sheets + Drop-Down with
# hyperlinks (fake schools, contoso links). Step 2 (here): add styles, CF and a DV list.
import subprocess
from openpyxl.formatting.rule import CellIsRule as _CIR
raw = os.path.join(HERE, "_audit_raw.xlsx")
subprocess.run(["node", os.path.join(HERE, "make-audit-fixture.cjs"), raw], check=True)
aw = openpyxl.load_workbook(raw)
for ws in aw:
    if ws.title == "Drop-Down": continue
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F6CBD")
    ws.column_dimensions["A"].width = 30
    ws.conditional_formatting.add("C2:C50", _CIR(operator="greaterThan", formula=["25"], fill=PatternFill("solid", fgColor="FDE7E9")))
    adv = DataValidation(type="list", formula1="'Drop-Down'!$A$2:$A$3", allow_blank=True); ws.add_data_validation(adv); adv.add("D2:D50")
aw.save(os.path.join(HERE, "SAMPLE_AUDIT_RICH.xlsx")); os.remove(raw)
print("wrote", os.path.join(HERE, "SAMPLE_AUDIT_RICH.xlsx"))
